"""AI 文件分析功能测试（核心层 + 本地模拟服务器 + 界面冒烟）。

覆盖范围：
* 核心层纯函数：collect_files / build_messages / _extract_json_array /
  parse_insights / AiClient 配置校验；
* 网络层端到端：用 ``http.server`` 在本地随机端口模拟 OpenAI 兼容接口，
  验证请求头、请求体、成功解析与各类错误分支（不访问真实外网）；
* 界面冒烟：AI 设置对话框保存往返、分析对话框打开即列清单（不自动请求）、
  未配置时点击“开始分析”不发请求、完整分析链路与结果导出。

运行方式::

    python -m unittest discover -s tests -v

说明：所有弹窗都被替换成假对象，测试不会真的弹窗、不会打开资源管理器。
"""

from __future__ import annotations

import importlib.util
import json
import logging
import os
import shutil
import sys
import tempfile
import threading
import time
import unittest
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, _ROOT)

from PySide6.QtWidgets import QApplication  # noqa: E402

from app.core import exporter  # noqa: E402
from app.core.ai_client import (  # noqa: E402
    AiAnalysisError,
    AiClient,
    FileEntry,
    build_messages,
    collect_files,
    entry_key,
    parse_insights,
    _extract_json_array,
)
from app.core.config import ConfigManager  # noqa: E402
from app.ui import ai_dialog as ad  # noqa: E402
from app.ui.ai_thread import AiAnalysisThread  # noqa: E402

logging.disable(logging.INFO)

HAS_OPENPYXL = importlib.util.find_spec("openpyxl") is not None


# ======================================================================
# 通用工具
# ======================================================================
def _app() -> QApplication:
    """进程内共享的 QApplication。"""
    return QApplication.instance() or QApplication([])


def _pump(app: QApplication, seconds: float) -> None:
    """让事件循环空转一段时间（等待界面刷新/信号派发）。"""
    deadline = time.monotonic() + seconds
    while time.monotonic() < deadline:
        app.processEvents()
        time.sleep(0.01)


def _wait_for(app: QApplication, predicate, timeout: float = 30.0) -> bool:
    """轮询等待条件成立，期间持续派发 Qt 事件。"""
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        app.processEvents()
        if predicate():
            return True
        time.sleep(0.02)
    return predicate()


def _make_entries() -> list:
    """构造一组固定的测试文件条目（不落盘，供提示词/解析测试用）。"""
    return [
        FileEntry("big.dll", r"C:\demo\big.dll", 4096, 1700000000.0),
        FileEntry("app.exe", r"C:\demo\app.exe", 2048, 1700000100.0),
        FileEntry("notes.txt", r"C:\demo\notes.txt", 512, 1700000200.0),
        FileEntry("settings.ini", r"C:\demo\settings.ini", 100, 0.0),
    ]


# ======================================================================
# 核心层单测（无 Qt、无网络）
# ======================================================================
class AiCoreTest(unittest.TestCase):
    """collect_files / build_messages / 解析 / 配置校验 的单元测试。"""

    def setUp(self) -> None:
        self.tmp = tempfile.mkdtemp(prefix="ai_core_")
        self.addCleanup(shutil.rmtree, self.tmp, ignore_errors=True)

    # -------------------------------------------------------- collect_files
    def test_collect_files_basic(self) -> None:
        """递归收集：包含子目录文件、rel_dir 正确、按大小降序、max_files 生效。"""
        # 3 个根文件 + 1 个子目录（子目录内文件现在也会被收集）
        with open(os.path.join(self.tmp, "a.txt"), "wb") as fp:
            fp.write(b"x" * 10)
        with open(os.path.join(self.tmp, "b.bin"), "wb") as fp:
            fp.write(b"y" * 2000)
        with open(os.path.join(self.tmp, "c.log"), "wb") as fp:
            fp.write(b"z" * 500)
        sub = os.path.join(self.tmp, "sub")
        os.mkdir(sub)
        with open(os.path.join(sub, "inner.txt"), "wb") as fp:
            fp.write(b"i" * 9999)

        entries = collect_files(self.tmp)
        self.assertEqual([e.name for e in entries],
                         ["inner.txt", "b.bin", "c.log", "a.txt"])  # 大 -> 小
        self.assertEqual(entries[0].rel_dir, "sub")          # 子目录文件
        self.assertEqual(entries[1].rel_dir, "")             # 根目录文件
        self.assertEqual(entries[1].path, os.path.join(self.tmp, "b.bin"))

        truncated = collect_files(self.tmp, max_files=2)
        self.assertEqual([e.name for e in truncated], ["inner.txt", "b.bin"])

    def test_collect_files_per_dir_limit(self) -> None:
        """每个子目录最多取 per_dir_limit 个文件（大文件优先）。"""
        d1 = os.path.join(self.tmp, "d1")
        d2 = os.path.join(self.tmp, "d2")
        os.mkdir(d1)
        os.mkdir(d2)
        for i, size in enumerate((100, 200, 300)):   # d1 里 3 个文件
            with open(os.path.join(d1, f"f{i}.bin"), "wb") as fp:
                fp.write(b"x" * size)
        for i, size in enumerate((60, 80)):          # d2 里 2 个文件
            with open(os.path.join(d2, f"g{i}.bin"), "wb") as fp:
                fp.write(b"x" * size)

        entries = collect_files(self.tmp, max_files=10, per_dir_limit=2)
        # d1 只取最大的 2 个（f2=300, f1=200），d2 取满 2 个
        self.assertEqual({e.name for e in entries},
                         {"f1.bin", "f2.bin", "g0.bin", "g1.bin"})
        self.assertEqual(entries[0].rel_dir, "d1")    # 全局最大是 d1/f2.bin
        self.assertEqual({e.rel_dir for e in entries}, {"d1", "d2"})

    def test_entry_key_normalization(self) -> None:
        """entry_key 归一化：大小写、反斜杠、重复目录前缀、./ 前缀。"""
        self.assertEqual(entry_key("", "Big.DLL"), "big.dll")
        self.assertEqual(entry_key("sub", "a.txt"), "sub/a.txt")
        self.assertEqual(entry_key("sub\\dir", "a.txt"), "sub/dir/a.txt")
        # AI 有时把目录重复拼进文件名（sub/sub/a.txt -> sub/a.txt）
        self.assertEqual(entry_key("sub", "sub/a.txt"), "sub/a.txt")
        self.assertEqual(entry_key("sub", "./a.txt"), "sub/a.txt")
        self.assertEqual(entry_key("sub", "a.txt") == entry_key("", "sub/a.txt"),
                         True)         # 两种写法等价

    def test_collect_files_edge_cases(self) -> None:
        """空目录返回空列表；无效路径抛 AiAnalysisError。"""
        self.assertEqual(collect_files(self.tmp), [])
        with self.assertRaises(AiAnalysisError):
            collect_files(os.path.join(self.tmp, "不存在的目录"))
        with self.assertRaises(AiAnalysisError):
            collect_files("")

    # ------------------------------------------------------ build_messages
    def test_build_messages_structure(self) -> None:
        """messages 结构：system + user 两条，且包含目录与文件清单。"""
        entries = _make_entries()
        messages = build_messages(entries, r"C:\demo")
        self.assertEqual(len(messages), 2)
        self.assertEqual(messages[0]["role"], "system")
        self.assertEqual(messages[1]["role"], "user")
        user_text = messages[1]["content"]
        for token in (r"C:\demo", "big.dll", "4096", "settings.ini",
                      "JSON", "owner", "purpose", "category",
                      "所在目录/文件名"):
            self.assertIn(token, user_text)

    def test_build_messages_with_rel_dir(self) -> None:
        """子目录文件在清单里以“目录/文件名”形式展示。"""
        entries = _make_entries()
        entries[0] = FileEntry("big.dll", r"C:\demo\sub\big.dll",
                               4096, 1700000000.0, rel_dir="sub")
        messages = build_messages(entries, r"C:\demo")
        user_text = messages[1]["content"]
        self.assertIn("sub/big.dll", user_text)
        self.assertIn("settings.ini", user_text)   # 根目录文件不带目录前缀

    # --------------------------------------------------- _extract_json_array
    def test_extract_json_array_variants(self) -> None:
        """能从纯 JSON / Markdown 围栏 / 夹杂文字中提取数组。"""
        self.assertEqual(_extract_json_array('[{"a":1}]'), '[{"a":1}]')
        fenced = "```json\n[{\"a\": 1}]\n```"
        self.assertEqual(_extract_json_array(fenced), '[{"a": 1}]')
        noisy = '分析结果如下：[{"a":1}] 以上仅供参考'
        self.assertEqual(_extract_json_array(noisy), '[{"a":1}]')
        for bad in ("", "   ", "没有数组哦", "[未闭合"):
            with self.assertRaises(AiAnalysisError):
                _extract_json_array(bad)

    # -------------------------------------------------------- parse_insights
    def test_parse_insights_alignment(self) -> None:
        """按文件名对齐：大小写不敏感、多余项忽略、顺序与清单一致。"""
        entries = _make_entries()
        text = json.dumps([
            {"name": "BIG.DLL", "owner": "微软", "purpose": "系统动态库",
             "category": "程序"},
            {"name": "notes.txt", "owner": "用户", "purpose": "笔记",
             "category": "文档"},
            {"name": "ghost.exe", "owner": "不存在", "purpose": "多余项",
             "category": "其他"},          # 不在清单中：应被忽略
        ], ensure_ascii=False)
        insights = parse_insights(text, entries)
        self.assertEqual(len(insights), len(entries))
        self.assertEqual([i.name for i in insights],
                         [e.name for e in entries])        # 顺序保持清单顺序
        self.assertTrue(insights[0].ok)
        self.assertEqual(insights[0].owner, "微软")          # 大小写不敏感命中
        self.assertTrue(insights[2].ok)                      # notes.txt 命中

        # app.exe 未被 AI 回答：生成 ok=False 占位
        self.assertFalse(insights[1].ok)
        self.assertEqual(insights[1].name, "app.exe")
        self.assertIn("未返回", insights[1].purpose)

    def test_parse_insights_same_name_dirs(self) -> None:
        """不同子目录同名文件：靠“目录/文件名”唯一键正确对齐。"""
        entries = [
            FileEntry("data.bin", r"C:\demo\p1\data.bin", 100, 0.0,
                      rel_dir="p1"),
            FileEntry("data.bin", r"C:\demo\p2\data.bin", 200, 0.0,
                      rel_dir="p2"),
            FileEntry("readme.txt", r"C:\demo\readme.txt", 50, 0.0),
        ]
        text = json.dumps([
            {"name": "p1/data.bin", "owner": "甲", "purpose": "项目一数据",
             "category": "配置"},
            {"name": "p2/data.bin", "owner": "乙", "purpose": "项目二数据",
             "category": "配置"},
            {"name": "readme.txt", "owner": "丙", "purpose": "说明",
             "category": "文档"},
        ], ensure_ascii=False)
        insights = parse_insights(text, entries)
        self.assertEqual(len(insights), 3)
        self.assertTrue(all(i.ok for i in insights))
        self.assertEqual([i.rel_dir for i in insights], ["p1", "p2", ""])
        self.assertEqual([i.owner for i in insights], ["甲", "乙", "丙"])

    def test_parse_insights_dedup_dir_prefix(self) -> None:
        """AI 把目录重复拼进 name（p1/p1/data.bin）也能归一化命中。"""
        entries = [FileEntry("data.bin", r"C:\demo\p1\data.bin", 100, 0.0,
                             rel_dir="p1")]
        text = json.dumps([
            {"name": "p1/p1/data.bin", "rel_dir": "p1", "owner": "甲",
             "purpose": "数据", "category": "配置"},
        ], ensure_ascii=False)
        insights = parse_insights(text, entries)
        self.assertTrue(insights[0].ok)
        self.assertEqual(insights[0].name, "data.bin")
        self.assertEqual(insights[0].rel_dir, "p1")

    def test_parse_insights_invalid(self) -> None:
        """非法 JSON / 非数组都要给出可读错误。"""
        entries = _make_entries()
        with self.assertRaises(AiAnalysisError):
            parse_insights("[{broken", entries)
        with self.assertRaises(AiAnalysisError):
            parse_insights('{"name": "big.dll"}', entries)   # 对象而非数组

    # ------------------------------------------------------------- AiClient
    def test_client_validate_and_endpoint(self) -> None:
        """缺配置时给出中文提示；endpoint 拼接正确且容忍尾部斜杠。"""
        client = AiClient("https://api.example.com/v1/", "sk-x", "m1")
        self.assertEqual(client.endpoint, "https://api.example.com/v1/chat/completions")
        client._validate()          # 配置齐全：不抛异常

        for kwargs, fragment in (
            ({"base_url": "", "api_key": "k", "model": "m"}, "API 地址"),
            ({"base_url": "https://a.com/v1", "api_key": "", "model": "m"}, "API Key"),
            ({"base_url": "https://a.com/v1", "api_key": "k", "model": ""}, "模型名称"),
        ):
            with self.assertRaises(AiAnalysisError) as ctx:
                AiClient(**kwargs)._validate()          # type: ignore[arg-type]
            self.assertIn(fragment, str(ctx.exception))


# ======================================================================
# 网络层端到端（本地模拟 OpenAI 兼容服务器）
# ======================================================================
class _StubApiHandler(BaseHTTPRequestHandler):
    """极简 OpenAI 兼容接口：记录请求，返回预设状态码与响应体。"""

    def do_POST(self) -> None:  # noqa: N802 - 基类命名约定
        length = int(self.headers.get("Content-Length") or 0)
        body = self.rfile.read(length) if length else b""
        self.server.requests.append({          # type: ignore[attr-defined]
            "path": self.path,
            "auth": self.headers.get("Authorization") or "",
            "body": body.decode("utf-8", "replace"),
        })
        code = getattr(self.server, "next_code", 200)
        payload = getattr(self.server, "next_payload", None)
        if payload is None:                    # 默认：成功的 chat 响应
            content = json.dumps([
                {"name": "big.dll", "owner": "微软 / Windows 系统组件",
                 "purpose": "Windows 系统动态链接库", "category": "程序"},
                {"name": "app.exe", "owner": "测试软件", "purpose": "演示程序",
                 "category": "程序"},
                {"name": "notes.txt", "owner": "用户本人", "purpose": "文本笔记",
                 "category": "文档"},
                {"name": "settings.ini", "owner": "测试软件", "purpose": "配置文件",
                 "category": "配置"},
                {"name": "sub/inner.bin", "owner": "测试软件",
                 "purpose": "子目录示例文件", "category": "程序"},
            ], ensure_ascii=False)
            payload = {"choices": [{"message": {"content": content}}]}
        data = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        self.send_response(code)
        self.send_header("Content-Type", "application/json")
        self.send_header("Content-Length", str(len(data)))
        self.end_headers()
        self.wfile.write(data)

    def log_message(self, *_args) -> None:     # 静默，避免污染测试输出
        pass


class AiClientNetworkTest(unittest.TestCase):
    """AiClient.chat 与模拟服务器的端到端测试（127.0.0.1 随机端口）。"""

    @classmethod
    def setUpClass(cls) -> None:
        cls.server = ThreadingHTTPServer(("127.0.0.1", 0), _StubApiHandler)
        cls.server.requests = []               # type: ignore[attr-defined]
        cls.thread = threading.Thread(
            target=cls.server.serve_forever, kwargs={"poll_interval": 0.05},
            daemon=True)
        cls.thread.start()
        cls.base_url = f"http://127.0.0.1:{cls.server.server_port}/v1"

    @classmethod
    def tearDownClass(cls) -> None:
        cls.server.shutdown()
        cls.server.server_close()
        cls.thread.join(timeout=3)

    def setUp(self) -> None:
        """每个用例都恢复默认响应，并准备一个配置齐全的客户端。"""
        self.server.next_code = 200            # type: ignore[attr-defined]
        self.server.next_payload = None        # type: ignore[attr-defined]
        self.server.requests.clear()           # type: ignore[attr-defined]
        self.client = AiClient(self.base_url, "sk-test-key", "stub-model",
                               timeout=10.0)

    def test_chat_success(self) -> None:
        """成功链路：请求头/请求体正确，回复文本可解析。"""
        entries = _make_entries()
        content = self.client.chat(build_messages(entries, r"C:\demo"))
        self.assertIn("big.dll", content)

        self.assertEqual(len(self.server.requests), 1)   # type: ignore[attr-defined]
        record = self.server.requests[0]                 # type: ignore[attr-defined]
        self.assertEqual(record["path"], "/v1/chat/completions")
        self.assertEqual(record["auth"], "Bearer sk-test-key")
        sent = json.loads(record["body"])
        self.assertEqual(sent["model"], "stub-model")
        self.assertFalse(sent["stream"])
        self.assertEqual(len(sent["messages"]), 2)

    def test_chat_http_401(self) -> None:
        """401 → 提示 API Key 无效。"""
        self.server.next_code = 401            # type: ignore[attr-defined]
        with self.assertRaises(AiAnalysisError) as ctx:
            self.client.chat(build_messages(_make_entries(), r"C:\demo"))
        self.assertIn("401", str(ctx.exception))
        self.assertIn("API Key 无效", str(ctx.exception))

    def test_chat_bad_response_format(self) -> None:
        """缺少 choices 字段 → 响应格式异常。"""
        self.server.next_payload = {"error": {"message": "nope"}}  # type: ignore[attr-defined]
        with self.assertRaises(AiAnalysisError) as ctx:
            self.client.chat(build_messages(_make_entries(), r"C:\demo"))
        self.assertIn("格式异常", str(ctx.exception))

    def test_chat_empty_content(self) -> None:
        """content 为空白 → 回复内容为空。"""
        self.server.next_payload = {"choices": [{"message": {"content": "  "}}]}  # type: ignore[attr-defined]
        with self.assertRaises(AiAnalysisError) as ctx:
            self.client.chat(build_messages(_make_entries(), r"C:\demo"))
        self.assertIn("为空", str(ctx.exception))

    def test_chat_rejects_missing_config(self) -> None:
        """配置缺失时不发任何网络请求。"""
        client = AiClient("", "", "")
        with self.assertRaises(AiAnalysisError):
            client.chat([{"role": "user", "content": "hi"}])
        self.assertEqual(len(self.server.requests), 0)   # type: ignore[attr-defined]


# ======================================================================
# 界面冒烟测试（离屏 Qt + 本地模拟服务器）
# ======================================================================
class StubMessageBox:
    """替代 ai_dialog 内的 QMessageBox：记录调用、exec 返回预设值。"""

    calls: list = []
    next_exec_result = 0            # 默认 0（等价于“否”），避免阻塞

    class StandardButton:
        Ok = 0x1
        Yes = 0x4
        No = 0x10

    class Icon:
        Warning = 1
        Information = 2
        Question = 3
        Critical = 4

    def __init__(self, _parent=None) -> None:
        self.text = ""

    def setWindowTitle(self, _title) -> None:
        pass

    def setText(self, text) -> None:
        self.text = str(text)

    def setIcon(self, _icon) -> None:
        pass

    def setStandardButtons(self, _buttons) -> None:
        pass

    def setDefaultButton(self, _button) -> None:
        pass

    def exec(self, *_args) -> int:
        StubMessageBox.calls.append(("exec", self.text))
        return StubMessageBox.next_exec_result

    @classmethod
    def information(cls, _parent, title, text):
        cls.calls.append(("information", title, text))

    @classmethod
    def warning(cls, _parent, title, text):
        cls.calls.append(("warning", title, text))

    @classmethod
    def critical(cls, _parent, title, text):
        cls.calls.append(("critical", title, text))


class FakeFileDialog:
    """替代 ai_dialog 内的 QFileDialog：返回预设保存路径。"""

    save_answer = ""

    @staticmethod
    def getSaveFileName(_parent, _title, suggested, _filter):
        return (FakeFileDialog.save_answer or suggested), ""


class AiDialogSmokeTest(unittest.TestCase):
    """AI 设置/分析对话框冒烟测试（不访问真实网络）。"""

    FILES = [("big.dll", 4096), ("app.exe", 2048),
             ("notes.txt", 512), ("settings.ini", 100)]
    SUB_FILES = [("inner.bin", 3000)]      # 子目录文件：验证递归收集在界面上生效

    @classmethod
    def setUpClass(cls) -> None:
        cls.app = _app()
        cls.tmp = tempfile.mkdtemp(prefix="ai_ui_")
        # 准备被分析目录：4 个根文件 + 1 个子目录（含 1 个文件） + 1 个空子目录
        cls.directory = os.path.join(cls.tmp, "target")
        os.mkdir(cls.directory)
        for name, size in cls.FILES:
            with open(os.path.join(cls.directory, name), "wb") as fp:
                fp.write(b"x" * size)
        os.mkdir(os.path.join(cls.directory, "sub"))
        for name, size in cls.SUB_FILES:
            with open(os.path.join(cls.directory, "sub", name), "wb") as fp:
                fp.write(b"x" * size)
        os.mkdir(os.path.join(cls.directory, "empty_sub"))
        # 备桩并替换 ai_dialog 模块内的符号
        cls._real = (ad.QMessageBox, ad.QFileDialog, ad.open_in_explorer)
        ad.QMessageBox = StubMessageBox          # type: ignore[misc]
        ad.QFileDialog = FakeFileDialog          # type: ignore[misc]
        ad.open_in_explorer = lambda *a, **k: None

    @classmethod
    def tearDownClass(cls) -> None:
        ad.QMessageBox, ad.QFileDialog, ad.open_in_explorer = cls._real
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def setUp(self) -> None:
        StubMessageBox.calls.clear()
        FakeFileDialog.save_answer = ""
        StubMessageBox.next_exec_result = 0
        # 每个用例使用独立配置：先删残留文件，避免上一用例保存的配置串扰
        self.cfg_path = os.path.join(self.tmp, "config.json")
        if os.path.exists(self.cfg_path):
            os.remove(self.cfg_path)
        self.cfg = ConfigManager(self.cfg_path)

    def tearDown(self) -> None:
        _pump(self.app, 0.1)

    def _open_dialog(self) -> "ad.AiAnalysisDialog":
        """打开对话框并等待后台收集线程结束（递归收集是异步的）。"""
        dialog = ad.AiAnalysisDialog(self.cfg, self.directory)
        self.addCleanup(dialog.close)
        self.assertTrue(_wait_for(
            self.app,
            lambda: (dialog._collector is None
                     or not dialog._collector.isRunning()),
            timeout=15.0), "文件清单收集线程未在超时内结束")
        return dialog

    # ------------------------------------------------------------ 设置对话框
    def test_settings_roundtrip(self) -> None:
        """填写参数并保存 → 配置正确写回（含类型转换）。"""
        dialog = ad.AiSettingsDialog(self.cfg)
        dialog.edit_base_url.setText("https://api.example.com/v1/")
        dialog.edit_api_key.setText("  sk-abc  ")          # 应被 strip
        dialog.edit_model.setText("demo-model")
        dialog.spin_batch.setValue(12)
        dialog.spin_max_files.setValue(150)
        dialog.accept()
        self.assertEqual(self.cfg.get("ai_base_url"), "https://api.example.com/v1")
        self.assertEqual(self.cfg.get("ai_api_key"), "sk-abc")
        self.assertEqual(self.cfg.get("ai_model"), "demo-model")
        self.assertEqual(self.cfg.get("ai_batch_size"), 12)
        self.assertEqual(self.cfg.get("ai_max_files"), 150)

    def test_settings_rejects_bad_url(self) -> None:
        """接口地址缺少 http(s) 前缀 → 提示且不保存。"""
        dialog = ad.AiSettingsDialog(self.cfg)
        dialog.edit_base_url.setText("api.example.com/v1")
        dialog.edit_api_key.setText("sk-abc")
        dialog.edit_model.setText("demo-model")
        dialog.accept()
        self.assertTrue(any(kind == "warning"
                            for kind, *_ in StubMessageBox.calls))
        self.assertEqual(self.cfg.get("ai_base_url"), "")   # 未保存

    # ------------------------------------------------------------ 分析对话框
    def test_dialog_open_lists_entries_without_request(self) -> None:
        """打开对话框只收集清单、不发起任何网络请求；按钮初始状态正确。"""
        dialog = self._open_dialog()
        self.assertEqual(len(dialog._entries), 5)    # 4 根文件 + sub/inner.bin
        # 子目录文件带上了所在目录信息
        self.assertTrue(any(e.name == "inner.bin" and e.rel_dir == "sub"
                            for e in dialog._entries))
        self.assertIn("5", dialog.lbl_status.text())
        self.assertFalse(dialog.btn_export_csv.isEnabled())
        self.assertFalse(dialog.btn_stop.isEnabled())
        self.assertEqual(self.server_requests(), [])        # 未发请求

    def test_dialog_empty_directory(self) -> None:
        """空目录：开始按钮禁用并给出提示。"""
        empty = os.path.join(self.tmp, "nothing")
        os.makedirs(empty, exist_ok=True)
        dialog = ad.AiAnalysisDialog(self.cfg, empty)
        self.addCleanup(dialog.close)
        self.assertTrue(_wait_for(
            self.app,
            lambda: (dialog._collector is None
                     or not dialog._collector.isRunning()),
            timeout=15.0))
        self.assertEqual(dialog._entries, [])
        self.assertFalse(dialog.btn_start.isEnabled())

    def test_start_without_config_does_not_request(self) -> None:
        """未配置接口时点击“开始分析”：弹提示、不建线程、不发请求。"""
        dialog = self._open_dialog()
        dialog.btn_start.click()
        self.assertIsNone(dialog._thread)
        self.assertTrue(StubMessageBox.calls)               # 出现了提示弹窗
        self.assertEqual(self.server_requests(), [])        # 未发请求

    # ------------------------------------------------------ 完整分析链路
    def server_requests(self) -> list:
        """读取当前类级模拟服务器的请求记录（未启动则返回空表）。"""
        return list(getattr(self, "_srv_requests", []))

    def _start_stub_server(self) -> None:
        """为完整链路测试启动一个独立的模拟服务器。"""
        self._server = ThreadingHTTPServer(("127.0.0.1", 0), _StubApiHandler)
        self._server.requests = []
        self._server_thread = threading.Thread(
            target=self._server.serve_forever, kwargs={"poll_interval": 0.05},
            daemon=True)
        self._server_thread.start()
        self._srv_requests = self._server.requests
        self.base_url = f"http://127.0.0.1:{self._server.server_port}/v1"
        self.addCleanup(self._stop_stub_server)

    def _stop_stub_server(self) -> None:
        try:
            self._server.shutdown()
            self._server.server_close()
            self._server_thread.join(timeout=3)
        except Exception:                                  # noqa: BLE001
            pass

    def test_full_analysis_and_export_csv(self) -> None:
        """配置完整 → 点击开始 → 模拟服务器返回 → 表格逐行显示 → 导出 CSV。"""
        self._start_stub_server()
        self.cfg.update({
            "ai_base_url": self.base_url,
            "ai_api_key": "sk-ui-test",
            "ai_model": "stub-model",
            "ai_batch_size": 2,                # 5 个文件 → 3 批请求
            "ai_max_files": 200,
        }, save=True)
        dialog = self._open_dialog()
        dialog.btn_start.click()

        total = len(self.FILES) + len(self.SUB_FILES)
        finished = _wait_for(
            self.app,
            lambda: (dialog._thread is not None
                     and not dialog._thread.isRunning()
                     and dialog.table.rowCount() == total),
            timeout=20.0)
        self.assertTrue(finished, "分析线程未在超时内完成")

        # 三个批次各请求一次；表格逐行包含全部文件与 AI 说明
        self.assertEqual(len(self._srv_requests), 3)
        self.assertEqual(dialog.table.rowCount(), total)
        names = {dialog.table.item(r, 0).text() for r in range(total)}
        self.assertEqual(names, {n for n, _ in self.FILES}
                         | {n for n, _ in self.SUB_FILES})
        self.assertTrue(dialog.table.item(0, 5).text())     # 用途说明非空
        # 子目录行显示了“所在目录”
        sub_rows = [r for r in range(total)
                    if dialog.table.item(r, 0).text() == "inner.bin"]
        self.assertEqual(dialog.table.item(sub_rows[0], 1).text(), "sub")
        self.assertTrue(dialog.btn_export_csv.isEnabled())
        self.assertTrue(dialog.btn_export_xlsx.isEnabled())
        self.assertTrue(dialog.btn_start.isEnabled())       # 按钮已恢复

        # 导出 CSV（假文件对话框返回临时路径）
        target = os.path.join(self.tmp, "报告.csv")
        FakeFileDialog.save_answer = target
        dialog.export_results("csv")
        self.assertTrue(os.path.isfile(target))
        with open(target, encoding="utf-8-sig") as fp:
            text = fp.read()
        self.assertIn("软件归属", text)
        self.assertIn("所在目录", text)
        self.assertIn("Windows 系统动态链接库", text)
        self.assertTrue(any(kind == "information"
                            for kind, *_ in StubMessageBox.calls))

    @unittest.skipUnless(HAS_OPENPYXL, "未安装 openpyxl，跳过 Excel 导出测试")
    def test_export_xlsx(self) -> None:
        """分析完成后导出 Excel 报告。"""
        self._start_stub_server()
        self.cfg.update({
            "ai_base_url": self.base_url,
            "ai_api_key": "sk-ui-test",
            "ai_model": "stub-model",
        }, save=True)
        dialog = self._open_dialog()
        dialog.btn_start.click()
        total = len(self.FILES) + len(self.SUB_FILES)
        self.assertTrue(_wait_for(
            self.app,
            lambda: (dialog._thread is not None
                     and not dialog._thread.isRunning()
                     and dialog.table.rowCount() == total),
            timeout=20.0))

        target = os.path.join(self.tmp, "报告.xlsx")
        FakeFileDialog.save_answer = target
        dialog.export_results("xlsx")
        self.assertTrue(os.path.isfile(target))

        # 用 openpyxl 读回校验标题与数据行
        from openpyxl import load_workbook
        workbook = load_workbook(target)
        sheet = workbook.active
        self.assertEqual(sheet.cell(row=1, column=1).value, "AI 文件分析报告")
        texts = [str(cell.value) for row in sheet.iter_rows()
                 for cell in row if cell.value]
        self.assertIn("软件归属", " ".join(texts))
        self.assertIn("所在目录", " ".join(texts))
        self.assertIn("big.dll", " ".join(texts))
        self.assertIn("inner.bin", " ".join(texts))


if __name__ == "__main__":
    unittest.main(verbosity=2)
