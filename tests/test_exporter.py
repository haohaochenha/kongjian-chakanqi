"""导出模块测试：CSV 编码/表头、XLSX 结构、以及目标路径校验。

运行方式::

    python -m unittest discover -s tests -v
"""

from __future__ import annotations

import csv
import importlib.util
import io
import logging
import os
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from app.core import exporter  # noqa: E402
from app.core.formatter import format_size  # noqa: E402

logging.disable(logging.INFO)

HAS_OPENPYXL = importlib.util.find_spec("openpyxl") is not None


def sample_rows() -> list:
    """构造两行示例数据（键与 exporter.COLUMNS 对齐）。"""
    return exporter.rows_from_records([
        {"name": "项目文件", "path": r"D:\空间查看器", "size": 3172, "percent": 0.8,
         "files": 4, "subdirs": 3, "direct_files": 1, "largest": 2048, "depth": 1,
         "modified": 1700000000.0},
        {"name": "空目录", "path": r"D:\empty", "size": 0, "percent": 0.0,
         "files": 0, "subdirs": 0, "direct_files": 0, "largest": 0, "depth": 0,
         "modified": 0},
    ])


class TestRowsFromRecords(unittest.TestCase):
    """导出行补齐逻辑。"""

    def test_size_text_generated(self) -> None:
        """未提供 size_text 时按字节自动生成。"""
        rows = sample_rows()
        self.assertEqual(rows[0]["size_text"], format_size(3172))
        self.assertIn("KB", rows[0]["size_text"])

    def test_modified_formatted_as_text(self) -> None:
        """时间戳转成可读文本，0 转为空串。"""
        rows = sample_rows()
        self.assertRegex(str(rows[0]["modified"]), r"^\d{4}-\d{2}-\d{2} \d{2}:\d{2}$")
        self.assertEqual(rows[1]["modified"], "")

    def test_explicit_size_text_kept(self) -> None:
        """已存在的 size_text 不被覆盖（界面可指定固定单位）。"""
        rows = exporter.rows_from_records([{"name": "x", "size": 1024, "size_text": "1.00 KB"}])
        self.assertEqual(rows[0]["size_text"], "1.00 KB")

    def test_missing_keys_tolerated(self) -> None:
        """缺字段的脏数据不抛异常。"""
        rows = exporter.rows_from_records([{"name": "只有名字"}])
        self.assertEqual(rows[0]["size_text"], "0 B")
        self.assertEqual(rows[0]["percent"], 0.0)


class TestCsvExport(unittest.TestCase):
    """CSV 导出（Excel 兼容性重点）。"""

    def setUp(self) -> None:
        self.dir = tempfile.mkdtemp(prefix="space-viewer-export-")
        self.addCleanup(self._cleanup)

    def _cleanup(self) -> None:
        import shutil
        shutil.rmtree(self.dir, ignore_errors=True)

    def test_suffix_and_bom(self) -> None:
        """自动补 .csv 后缀，且带 UTF-8 BOM（Excel 打开中文不乱码）。"""
        target = os.path.join(self.dir, "报告")
        path = exporter.export_csv(target, sample_rows(), {"扫描根目录": "D:\\"})
        self.assertTrue(path.endswith(".csv"))
        with open(path, "rb") as fp:
            self.assertEqual(fp.read(3), b"\xef\xbb\xbf")

    def test_header_and_rows(self) -> None:
        """表头顺序与 COLUMNS 一致，数据行数正确。"""
        path = exporter.export_csv(os.path.join(self.dir, "a.csv"), sample_rows())
        with open(path, "r", encoding="utf-8-sig", newline="") as fp:
            table = list(csv.reader(fp))
        self.assertEqual(table[0], [title for title, *_ in exporter.COLUMNS])
        self.assertEqual(len(table), 1 + len(sample_rows()))
        self.assertEqual(table[1][0], "项目文件")
        # 逗号/引号安全转义
        tricky = exporter.rows_from_records([{"name": '含,逗号与"引号"', "size": 1}])
        path2 = exporter.export_csv(os.path.join(self.dir, "b.csv"), tricky)
        with open(path2, "r", encoding="utf-8-sig", newline="") as fp:
            back = list(csv.reader(fp))
        self.assertEqual(back[1][0], '含,逗号与"引号"')

    def test_meta_written_as_comment(self) -> None:
        """meta 以 # 前缀写在表头之前。"""
        path = exporter.export_csv(os.path.join(self.dir, "c.csv"), sample_rows(),
                                   {"标题": "文件夹排行", "总占用": "3.10 KB"})
        with open(path, "r", encoding="utf-8-sig", newline="") as fp:
            first = next(iter(csv.reader(fp)))
        self.assertTrue(first[0].startswith("#"))
        self.assertEqual(first[1], "文件夹排行")

    def test_target_errors(self) -> None:
        """目录不存在 / 目标是目录 时给出明确异常。"""
        with self.assertRaises(FileNotFoundError):
            exporter.export_csv(os.path.join(self.dir, "nope", "x.csv"), sample_rows())
        fake = os.path.join(self.dir, "isdir.csv")
        os.makedirs(fake, exist_ok=True)
        with self.assertRaises(IsADirectoryError):
            exporter.export_csv(fake, sample_rows())


@unittest.skipUnless(HAS_OPENPYXL, "未安装 openpyxl，跳过 Excel 导出测试")
class TestXlsxExport(unittest.TestCase):
    """XLSX 导出（回读校验结构与样式）。"""

    def setUp(self) -> None:
        self.dir = tempfile.mkdtemp(prefix="space-viewer-xlsx-")
        self.addCleanup(self._cleanup)

    def _cleanup(self) -> None:
        import shutil
        shutil.rmtree(self.dir, ignore_errors=True)

    def _load(self, path: str):
        from openpyxl import load_workbook
        return load_workbook(path)

    def test_structure(self) -> None:
        """标题、表头、数据区、冻结窗格与自动筛选均存在。"""
        from openpyxl.utils import get_column_letter

        path = exporter.export_xlsx(os.path.join(self.dir, "report"), sample_rows(),
                                    exporter.build_meta("D:\\", 3172, 4, 4, 1.2,
                                                        "folders", 2))
        self.assertTrue(path.endswith(".xlsx"))
        workbook = self._load(path)
        sheet = workbook.active
        self.assertEqual(sheet.title, "文件夹排行")
        self.assertEqual(sheet.cell(row=1, column=1).value, "文件夹排行")

        header_row = None
        titles = [title for title, *_ in exporter.COLUMNS]
        for r in range(1, 12):
            if sheet.cell(row=r, column=1).value == titles[0]:
                header_row = r
                break
        self.assertIsNotNone(header_row, "未找到表头行")
        self.assertEqual([sheet.cell(row=header_row, column=c).value
                          for c in range(1, len(titles) + 1)], titles)
        # 数据区
        self.assertEqual(sheet.cell(row=header_row + 1, column=1).value, "项目文件")
        self.assertEqual(sheet.cell(row=header_row + 2, column=1).value, "空目录")
        self.assertEqual(sheet.cell(row=header_row + 1, column=4).value, 3172)
        # 修改时间列：有时间戳的行是文本，无时间戳（0）的行必须是空串而不是数字 0
        self.assertRegex(str(sheet.cell(row=header_row + 1, column=11).value),
                         r"^\d{4}-\d{2}-\d{2} \d{2}:\d{2}")
        self.assertIn(sheet.cell(row=header_row + 2, column=11).value, ("", None))
        # 冻结与筛选
        self.assertEqual(sheet.freeze_panes, f"A{header_row + 1}")
        self.assertTrue(str(sheet.auto_filter.ref).startswith(f"A{header_row}:"))
        self.assertIn(f"{get_column_letter(len(titles))}{header_row + 2}",
                      str(sheet.auto_filter.ref))
        workbook.close()

    def test_extension_report_xlsx(self) -> None:
        """文件类型分布报告可正常生成。"""
        path = exporter.export_extension_report(
            os.path.join(self.dir, "ext.xlsx"),
            {".bin": 3072, ".txt": 100}, {".bin": 2, ".txt": 1}, 3172)
        workbook = self._load(path)
        sheet = workbook.active
        self.assertEqual(sheet.cell(row=1, column=1).value, "文件类型占用分布")
        workbook.close()

    def test_empty_rows_still_writes_header(self) -> None:
        """空结果集也应生成可打开的文件。"""
        path = exporter.export_xlsx(os.path.join(self.dir, "empty.xlsx"), [])
        workbook = self._load(path)
        self.assertIs(workbook.active.auto_filter.ref, None)
        workbook.close()


class TestBuildMeta(unittest.TestCase):
    """元信息文本组装。"""

    def test_fields(self) -> None:
        meta = exporter.build_meta(r"D:\data", 3172, 4, 4, 65.0, "folders", 100)
        self.assertEqual(meta["标题"], exporter.KIND_LABELS["folders"])
        self.assertEqual(meta["扫描根目录"], r"D:\data")
        self.assertEqual(meta["导出行数"], "100")
        self.assertIn("分", meta["扫描耗时"])   # 65 秒 → “1 分 5 秒”
        self.assertRegex(meta["导出时间"], r"^\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2}$")

    def test_unknown_kind_fallback(self) -> None:
        """未知类型直接沿用原字符串。"""
        self.assertEqual(exporter.build_meta("", 0, 0, 0, 0, "custom", 0)["标题"], "custom")


class TestStreamFallback(unittest.TestCase):
    """确保 csv 模块在文本流写入模式下可用（防止误改成二进制模式）。"""

    def test_write_to_text_io(self) -> None:
        """csv 写入内容可被标准库解析回来。"""
        buffer = io.StringIO()
        writer = csv.writer(buffer)
        writer.writerow([title for title, *_ in exporter.COLUMNS])
        buffer.seek(0)
        rows = list(csv.reader(buffer))
        self.assertEqual(len(rows[0]), len(exporter.COLUMNS))


if __name__ == "__main__":
    unittest.main(verbosity=2)
